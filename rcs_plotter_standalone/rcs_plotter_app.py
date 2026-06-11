#!/usr/bin/env python
# -*- coding: utf-8 -*-

from __future__ import annotations

import csv
import hashlib
import math
import sys
import tempfile
from dataclasses import dataclass
from collections import defaultdict
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

import numpy as np
from PyQt5 import QtCore, QtWidgets

import matplotlib

matplotlib.use("Qt5Agg")
from matplotlib.backends.backend_qt5agg import FigureCanvasQTAgg as FigureCanvas
from matplotlib.backends.backend_qt5agg import NavigationToolbar2QT as NavigationToolbar
from matplotlib.figure import Figure

try:
    from scipy.signal import savgol_filter as _savgol_filter
except Exception:  # pragma: no cover
    _savgol_filter = None

try:
    from scipy.interpolate import make_interp_spline as _make_interp_spline
except Exception:  # pragma: no cover
    _make_interp_spline = None

try:
    import openpyxl
    from openpyxl.chart import LineChart, Reference
    from openpyxl.drawing.image import Image as OpenpyxlImage
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except Exception:  # pragma: no cover
    openpyxl = None

import rcs_reference_data


APP_TITLE = "RCS 数据绘图工具"
RCS_MAX_DISTANCE_M = 50.0
RCS_STRAIGHT_X_MIN_M = 4.0
RCS_STRAIGHT_X_MAX_M = 50.0
RCS_FIT_GRID_STEP_M = 0.1
RCS_BIN_INCOHERENT_SUM_MAX_TIME_SPAN_S = 0.08
CLUSTER_RCS_PARSE_MAX_DIST_STEP_M = 5.0
CLUSTER_RCS_PARSE_TIME_GAP_RESET_S = 0.35
MAX_CLUSTERS = 20
ORBIT_POLAR_FIXED_RADIUS = 1.0
ORBIT_LINE_SMOOTH_POINTS = 360


@dataclass
class CurvePoint:
    t: float
    x_raw: float
    y_raw: float
    r_raw: float
    x: float
    y: float
    rcs_raw: float
    rcs_filt: float


@dataclass
class LoadedCurve:
    file_path: Path
    display_name: str
    segments: List[List[CurvePoint]]
    fitted: Tuple[np.ndarray, np.ndarray]

    @property
    def point_count(self) -> int:
        return sum(len(seg) for seg in self.segments)


@dataclass
class PointCloudPoint:
    file_path: Path
    source_index: int
    idx: int
    x: float
    y: float
    rcs_raw: float
    raw_x: float
    raw_y: float
    center_x: float
    center_y: float
    slot: int
    range_m: float
    viewing_angle: float


def _apply_matplotlib_font() -> None:
    matplotlib.rcParams["axes.unicode_minus"] = False
    preferred = [
        "Microsoft YaHei",
        "SimHei",
        "Noto Sans CJK SC",
        "Arial Unicode MS",
        "DejaVu Sans",
    ]
    matplotlib.rcParams["font.family"] = "sans-serif"
    matplotlib.rcParams["font.sans-serif"] = preferred


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _cell_float(cell: object) -> float:
    s = _cell_text(cell)
    if not s or s.upper() == "NAN":
        return float("nan")
    return float(s)


def _read_table_rows(path: Path) -> List[List[str]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as f:
            return [[_cell_text(c) for c in row] for row in csv.reader(f)]
    if suffix in {".xlsx", ".xlsm"}:
        if openpyxl is None:
            raise ValueError("当前环境缺少 openpyxl，不能读取 Excel 文件")
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb.active
            return [[_cell_text(c) for c in row] for row in ws.iter_rows(values_only=True)]
        finally:
            wb.close()
    raise ValueError("仅支持 CSV、XLSX、XLSM 表格文件")


def combine_rcs_db_incoherent_sum(values: Sequence[float]) -> Optional[float]:
    valid = []
    for value in values:
        try:
            v = float(value)
        except Exception:
            continue
        if math.isfinite(v):
            valid.append(v)
    if not valid:
        return None
    if len(valid) == 1:
        return float(valid[0])
    p_sum = sum(10.0 ** (v / 10.0) for v in valid)
    if p_sum <= 0.0 or not math.isfinite(p_sum):
        return None
    return float(10.0 * math.log10(p_sum))


def _smooth_rcs_series(y: np.ndarray) -> np.ndarray:
    y = np.asarray(y, dtype=float)
    if y.size <= 3:
        return y.copy()
    if _savgol_filter is not None:
        max_odd = y.size if y.size % 2 == 1 else y.size - 1
        win = min(39, max_odd)
        if win >= 3:
            if win % 2 == 0:
                win -= 1
            return np.asarray(_savgol_filter(y, win, 1, mode="nearest"), dtype=float)

    # Fallback: small median-like clipping followed by EMA.
    out = y.copy()
    alpha = 0.5
    for i in range(1, out.size):
        if abs(float(out[i]) - float(out[i - 1])) > 10.0:
            out[i] = out[i - 1]
        out[i] = alpha * out[i] + (1.0 - alpha) * out[i - 1]
    return out


def _cluster_point_from_slot(t_val: float, seg_key: int, x_raw: float, y_raw: float, rcs_val: float) -> Tuple[int, CurvePoint]:
    r_slant = math.hypot(float(x_raw), float(y_raw))
    return (
        int(seg_key),
        CurvePoint(
            t=float(t_val),
            x_raw=float(x_raw),
            y_raw=float(y_raw),
            r_raw=float(r_slant),
            x=float(x_raw),
            y=float(y_raw),
            rcs_raw=float(rcs_val),
            rcs_filt=float(rcs_val),
        ),
    )


def _parse_cluster_slot(row: List[str], col: Dict[str, int], slot_index: int) -> Optional[Tuple[float, float, float]]:
    keys = (f"DX{slot_index:02d}", f"DY{slot_index:02d}", f"RCS{slot_index:02d}")
    if any(k not in col for k in keys):
        return None
    try:
        dx = _cell_float(row[col[keys[0]]])
        dy = _cell_float(row[col[keys[1]]])
        rcs = _cell_float(row[col[keys[2]]])
    except Exception:
        return None
    if all(math.isfinite(v) for v in (dx, dy, rcs)):
        return float(dx), float(dy), float(rcs)
    return None


def _cluster_rowwise_merged_point(row: List[str], col: Dict[str, int], t_val: float, seg_key: int) -> List[Tuple[int, CurvePoint]]:
    s0 = _parse_cluster_slot(row, col, 0)
    s1 = _parse_cluster_slot(row, col, 1)
    if s0 is not None and s1 is not None:
        x0, y0, r0 = s0
        _x1, _y1, r1 = s1
        rcs_m = combine_rcs_db_incoherent_sum([r0, r1])
        if rcs_m is None:
            rcs_m = r0
        return [_cluster_point_from_slot(t_val, seg_key, x0, y0, rcs_m)]
    if s0 is not None:
        return [_cluster_point_from_slot(t_val, seg_key, s0[0], s0[1], s0[2])]
    if s1 is not None:
        return [_cluster_point_from_slot(t_val, seg_key, s1[0], s1[1], s1[2])]
    return []


def _point_with_time(p: CurvePoint, t_new: float) -> CurvePoint:
    return CurvePoint(
        t=float(t_new),
        x_raw=p.x_raw,
        y_raw=p.y_raw,
        r_raw=p.r_raw,
        x=p.x,
        y=p.y,
        rcs_raw=p.rcs_raw,
        rcs_filt=p.rcs_filt,
    )


def _finalize_cluster_points(points: List[CurvePoint]) -> List[CurvePoint]:
    if not points:
        return []
    pts = sorted(points, key=lambda p: (float(p.t), float(p.x_raw)))
    filtered = [pts[0]]
    last = pts[0]
    for p in pts[1:]:
        dt = float(p.t) - float(last.t)
        if not math.isfinite(dt):
            continue
        if abs(dt) < 1e-9 or dt >= CLUSTER_RCS_PARSE_TIME_GAP_RESET_S:
            filtered.append(p)
            last = p
            continue
        dist = math.hypot(float(p.x_raw) - float(last.x_raw), float(p.y_raw) - float(last.y_raw))
        if dist <= CLUSTER_RCS_PARSE_MAX_DIST_STEP_M:
            filtered.append(p)
            last = p
    t0 = float(filtered[0].t)
    return [_point_with_time(p, float(p.t) - t0) for p in filtered]


def parse_cluster_rcs_table(path: Path) -> Tuple[List[List[CurvePoint]], List[int]]:
    rows = _read_table_rows(path)
    header_index = -1
    col: Dict[str, int] = {}
    for i, row in enumerate(rows):
        if row and _cell_text(row[0]) == "Time" and any(_cell_text(c) == "DX00" for c in row):
            header = [_cell_text(c) for c in row]
            col = {name: idx for idx, name in enumerate(header) if name}
            header_index = i
            break
    if header_index < 0:
        raise ValueError("未找到 Cluster RCS 表头行（需要包含 Time 和 DX00）")
    for required in ("Time", "DX00", "DY00", "RCS00"):
        if required not in col:
            raise ValueError(f"Cluster RCS 表缺少列：{required}")

    max_ix = max(col.values())
    use_seg = "SegIdx" in col
    tagged: List[Tuple[int, CurvePoint]] = []
    for raw_row in rows[header_index + 1 :]:
        if not raw_row or not any(_cell_text(c) for c in raw_row):
            continue
        row = list(raw_row)
        while len(row) <= max_ix:
            row.append("")
        try:
            t_val = _cell_float(row[col["Time"]])
        except Exception:
            continue
        if not math.isfinite(t_val):
            continue
        seg_key = 0
        if use_seg:
            try:
                seg_key = int(round(_cell_float(row[col["SegIdx"]])))
            except Exception:
                seg_key = 0
        tagged.extend(_cluster_rowwise_merged_point(row, col, t_val, seg_key))

    if not tagged:
        raise ValueError("表格中没有可用的 DX/DY/RCS 数据点")

    if use_seg:
        by_seg: Dict[int, List[CurvePoint]] = {}
        for seg_key, point in tagged:
            by_seg.setdefault(int(seg_key), []).append(point)
        segments = [_finalize_cluster_points(by_seg[k]) for k in sorted(by_seg)]
        segments = [seg for seg in segments if seg]
        return segments, list(range(1, len(segments) + 1))

    points = [p for _, p in tagged]
    return [_finalize_cluster_points(points)], [1]


def parse_frame_aligned_point_cloud(
    path: Path,
    *,
    source_index: int = 1,
    start_idx: int = 1,
) -> Tuple[List[PointCloudPoint], int]:
    """Align every frame to DX00/DY00 and expand all valid cluster slots."""
    rows = _read_table_rows(path)
    header_index = -1
    col: Dict[str, int] = {}
    for i, row in enumerate(rows):
        if row and _cell_text(row[0]) == "Time" and any(_cell_text(c) == "DX00" for c in row):
            header = [_cell_text(c) for c in row]
            col = {name: idx for idx, name in enumerate(header) if name}
            header_index = i
            break
    if header_index < 0:
        raise ValueError("未找到点云数据表头，需要 Time、R、DX00、DY00、RCS00")
    for required in ("R", "DX00", "DY00", "RCS00"):
        if required not in col:
            raise ValueError(f"点云数据缺少列: {required}")

    max_ix = max(col.values())
    points: List[PointCloudPoint] = []
    idx = int(start_idx)
    for raw_row in rows[header_index + 1 :]:
        if not raw_row or not any(_cell_text(c) for c in raw_row):
            continue
        row = list(raw_row)
        while len(row) <= max_ix:
            row.append("")
        try:
            range_m = _cell_float(row[col["R"]])
            center_dx = _cell_float(row[col["DX00"]])
            center_dy = _cell_float(row[col["DY00"]])
        except Exception:
            continue
        if not all(math.isfinite(v) for v in (range_m, center_dx, center_dy)):
            continue

        center_x = float(center_dx) - float(range_m)
        center_y = float(center_dy)
        viewing_angle = float("nan")
        if "ViewAng" in col:
            try:
                viewing_angle = _cell_float(row[col["ViewAng"]])
            except Exception:
                viewing_angle = float("nan")

        for slot in range(MAX_CLUSTERS):
            parsed = _parse_cluster_slot(row, col, slot)
            if parsed is None:
                continue
            dx, dy, rcs = parsed
            if float(rcs) == -1.0:
                continue
            raw_x = float(dx) - float(range_m)
            raw_y = float(dy)
            points.append(
                PointCloudPoint(
                    file_path=path,
                    source_index=int(source_index),
                    idx=idx,
                    x=raw_x - center_x,
                    y=raw_y - center_y,
                    rcs_raw=float(rcs),
                    raw_x=raw_x,
                    raw_y=raw_y,
                    center_x=center_x,
                    center_y=center_y,
                    slot=int(slot),
                    range_m=float(range_m),
                    viewing_angle=float(viewing_angle),
                )
            )
            idx += 1

    if not points:
        raise ValueError("表格中没有可用于点云图的有效簇数据")
    return points, idx


def is_orbit_rcs_table(path: Path) -> bool:
    if "__orbit_rcs__" in path.name:
        return True
    try:
        rows = _read_table_rows(path)
    except Exception:
        return False
    if not rows:
        return False
    header = {str(c).strip().lower() for c in rows[0] if str(c).strip()}
    required = {"t", "theta_rad", "rcs_raw", "rcs_filt", "oid"}
    has_xy = ("x_m" in header or "x" in header) and ("y_m" in header or "y" in header)
    return required.issubset(header) and has_xy


def parse_orbit_rcs_table(path: Path) -> List[Dict[str, float]]:
    rows = _read_table_rows(path)
    if not rows:
        raise ValueError("圆周 RCS 表格为空")
    header = [_cell_text(c) for c in rows[0]]
    lower = {h.lower(): i for i, h in enumerate(header) if h}

    def col(*names: str) -> int:
        for name in names:
            if name in lower:
                return lower[name]
        raise ValueError(f"圆周 RCS 表缺少列：{'/'.join(names)}")

    c_t = col("t", "time_s", "time")
    c_rf = col("rcs_filt", "rcs_filtered")
    c_rr = lower.get("rcs_raw", c_rf)
    c_th = lower.get("theta_rad")
    c_x = lower.get("x_m", lower.get("x"))
    c_y = lower.get("y_m", lower.get("y"))
    c_oid = lower.get("oid", lower.get("cluster_id", -1))
    out: List[Dict[str, float]] = []
    max_ix = max(i for i in (c_t, c_rf, c_rr, c_th or 0, c_x or 0, c_y or 0, c_oid if c_oid >= 0 else 0))
    for raw in rows[1:]:
        row = list(raw)
        while len(row) <= max_ix:
            row.append("")
        try:
            item = {
                "t": _cell_float(row[c_t]),
                "rcs_raw": _cell_float(row[c_rr]),
                "rcs_filt": _cell_float(row[c_rf]),
                "theta_rad": _cell_float(row[c_th]) if c_th is not None else float("nan"),
                "x": _cell_float(row[c_x]) if c_x is not None else float("nan"),
                "y": _cell_float(row[c_y]) if c_y is not None else float("nan"),
                "oid": int(round(_cell_float(row[c_oid]))) if c_oid >= 0 else -1,
            }
        except Exception:
            continue
        if math.isfinite(item["t"]) and math.isfinite(item["rcs_filt"]):
            out.append(item)
    if len(out) < 2:
        raise ValueError("圆周 RCS 有效点不足")
    return sorted(out, key=lambda r: float(r["t"]))


def orbit_rows_from_cluster_segments(segments: Iterable[List[CurvePoint]]) -> List[Dict[str, float]]:
    rows: List[Dict[str, float]] = []
    for seg in segments:
        for p in seg:
            rows.append({"t": float(p.t), "rcs_filt": float(p.rcs_filt)})
    return sorted(rows, key=lambda r: float(r["t"]))


def build_orbit_series(rows: List[Dict[str, float]]) -> Dict[str, np.ndarray]:
    rows = sorted(rows, key=lambda r: float(r["t"]))
    times = np.asarray([float(r["t"]) for r in rows], dtype=float)
    values = np.asarray([float(r["rcs_filt"]) for r in rows], dtype=float)
    valid = np.isfinite(times) & np.isfinite(values)
    times = times[valid]
    values = values[valid]
    if times.size < 2:
        raise ValueError("圆周 RCS 有效点不足")
    span = float(times[-1] - times[0])
    if span <= 1e-12:
        angles = np.linspace(0.0, 2.0 * math.pi, times.size, endpoint=True)
    else:
        angles = (times - times[0]) / span * (2.0 * math.pi)
    return {
        "angles": angles,
        "values": values,
    }


def smooth_orbit_line(angles_rad: np.ndarray, values: np.ndarray) -> Tuple[np.ndarray, np.ndarray]:
    angles_rad = np.asarray(angles_rad, dtype=float)
    values = np.asarray(values, dtype=float)
    n = min(angles_rad.size, values.size)
    if n < 4 or _make_interp_spline is None:
        return angles_rad[:n], values[:n]

    angles_deg = np.mod(np.rad2deg(angles_rad[:n]), 360.0)
    rcs_values = values[:n]
    order = np.argsort(angles_deg)
    angles_deg = angles_deg[order]
    rcs_values = rcs_values[order]
    unique_angles, unique_indices = np.unique(angles_deg, return_index=True)
    angles_deg = unique_angles
    rcs_values = rcs_values[unique_indices]
    if angles_deg.size < 4:
        return angles_rad[:n], values[:n]

    try:
        angles_ext = np.concatenate([angles_deg - 360.0, angles_deg, angles_deg + 360.0])
        values_ext = np.concatenate([rcs_values, rcs_values, rcs_values])
        target_deg = np.linspace(0.0, 360.0, ORBIT_LINE_SMOOTH_POINTS, endpoint=False)
        k = min(3, len(angles_ext) - 1)
        spline = _make_interp_spline(angles_ext, values_ext, k=k)
        return np.deg2rad(target_deg), np.asarray(spline(target_deg), dtype=float)
    except Exception:
        return angles_rad[:n], values[:n]


def single_measurement_curve(points: List[CurvePoint]) -> Tuple[np.ndarray, np.ndarray]:
    if not points:
        return np.asarray([], dtype=float), np.asarray([], dtype=float)
    xs = np.asarray([float(p.x) for p in points], dtype=float)
    ys = np.asarray([float(p.rcs_filt) for p in points], dtype=float)
    ts = np.asarray([float(p.t) for p in points], dtype=float)
    mask = np.isfinite(xs) & np.isfinite(ys) & np.isfinite(ts)
    mask &= (xs >= RCS_STRAIGHT_X_MIN_M) & (xs <= RCS_STRAIGHT_X_MAX_M)
    xs, ys, ts = xs[mask], ys[mask], ts[mask]
    if xs.size == 0:
        return np.asarray([], dtype=float), np.asarray([], dtype=float)

    order = np.argsort(xs)
    xs, ys, ts = xs[order], ys[order], ts[order]
    if xs.size >= 2:
        keep = np.ones(xs.size, dtype=bool)
        last = 0
        for i in range(1, xs.size):
            if abs(float(ys[i]) - float(ys[last])) > 10.0:
                keep[i] = False
            else:
                last = i
        xs, ys, ts = xs[keep], ys[keep], ts[keep]
    if xs.size == 0:
        return np.asarray([], dtype=float), np.asarray([], dtype=float)

    bin_ids = np.floor(xs / RCS_FIT_GRID_STEP_M).astype(np.int64)
    xb: List[float] = []
    yb: List[float] = []
    for bid in np.unique(bin_ids):
        m = bin_ids == bid
        xb.append(float(np.mean(xs[m])))
        ys_b = ys[m]
        ts_b = ts[m]
        if ys_b.size <= 1:
            yb.append(float(ys_b[0]))
        elif float(np.max(ts_b) - np.min(ts_b)) <= RCS_BIN_INCOHERENT_SUM_MAX_TIME_SPAN_S:
            yb.append(float(combine_rcs_db_incoherent_sum(ys_b.tolist()) or np.mean(ys_b)))
        else:
            yb.append(float(np.mean(ys_b)))

    x_out = np.asarray(xb, dtype=float)
    y_out = np.asarray(yb, dtype=float)
    order = np.argsort(x_out)
    return x_out[order], _smooth_rcs_series(y_out[order])


def fuse_measurement_curves(curves: List[Tuple[np.ndarray, np.ndarray]]) -> Tuple[np.ndarray, np.ndarray]:
    valid: List[Tuple[np.ndarray, np.ndarray]] = []
    for xs, ys in curves:
        x = np.asarray(xs, dtype=float)
        y = np.asarray(ys, dtype=float)
        n = min(int(x.size), int(y.size))
        if n <= 0:
            continue
        x = x[:n]
        y = y[:n]
        mask = np.isfinite(x) & np.isfinite(y)
        if np.any(mask):
            valid.append((x[mask], y[mask]))
    if not valid:
        return np.asarray([], dtype=float), np.asarray([], dtype=float)
    if len(valid) == 1:
        return valid[0][0].copy(), valid[0][1].copy()

    bin_m = max(float(RCS_FIT_GRID_STEP_M), 1e-4)
    threshold_db = 10.0
    bin_to_xs: Dict[int, List[float]] = {}
    bin_to_ys: Dict[int, List[float]] = {}
    for xb, yb in valid:
        bids = np.floor(xb / bin_m).astype(np.int64)
        for bid, xk, yk in zip(bids.tolist(), xb.tolist(), yb.tolist()):
            bin_to_xs.setdefault(int(bid), []).append(float(xk))
            bin_to_ys.setdefault(int(bid), []).append(float(yk))

    x_all: List[float] = []
    y_all: List[float] = []
    for bid in sorted(bin_to_ys.keys()):
        ys_list = bin_to_ys.get(int(bid), [])
        xs_list = bin_to_xs.get(int(bid), [])
        if not ys_list or not xs_list:
            continue
        ys_arr = np.asarray(ys_list, dtype=float)
        xs_arr = np.asarray(xs_list, dtype=float)
        median = float(np.median(ys_arr))
        keep = np.abs(ys_arr - median) <= threshold_db
        if not np.any(keep):
            continue
        x_all.append(float(np.mean(xs_arr[keep])))
        y_all.append(float(np.mean(ys_arr[keep])))

    x_out = np.asarray(x_all, dtype=float)
    y_out = np.asarray(y_all, dtype=float)
    if x_out.size == 0:
        return x_out, y_out
    order = np.argsort(x_out)
    return x_out[order], _smooth_rcs_series(y_out[order])


def fit_segments(segments: List[List[CurvePoint]]) -> Tuple[np.ndarray, np.ndarray]:
    return fuse_measurement_curves([single_measurement_curve(seg) for seg in segments if seg])


def fit_loaded_curves_combined(curves: List[LoadedCurve]) -> Tuple[np.ndarray, np.ndarray]:
    merged_segments: List[List[CurvePoint]] = []
    for curve in curves:
        merged_segments.extend([list(seg) for seg in curve.segments if seg])
    return fit_segments(merged_segments)


def load_distance_curve(path: Path, display_name: Optional[str] = None) -> LoadedCurve:
    segments, _ = parse_cluster_rcs_table(path)
    fitted = fit_segments(segments)
    return LoadedCurve(path, display_name or path.stem, segments, fitted)


def _filtered_rcs_csv_path(raw_path: Path) -> Path:
    return raw_path.with_name(f"{raw_path.stem}_Filtered.csv")


def _combined_rcs_csv_path(raw_paths: List[Path]) -> Optional[Path]:
    paths = [Path(p) for p in raw_paths if str(p).strip()]
    if not paths:
        return None
    first = paths[0]
    if len(paths) == 1:
        return first.with_name(f"{first.stem}_combined.csv")
    source_key = "\n".join(str(p.resolve()) if p.exists() else str(p) for p in paths)
    digest = hashlib.sha1(source_key.encode("utf-8")).hexdigest()[:8]
    return first.with_name(f"{first.stem}_combined_{len(paths)}files_{digest}.csv")


def _cluster_slot_dx_rcs(row: List[str], col: Dict[str, int], slot_index: int) -> Optional[Tuple[float, float]]:
    dx_key = f"DX{slot_index:02d}"
    rcs_key = f"RCS{slot_index:02d}"
    if dx_key not in col or rcs_key not in col:
        return None
    try:
        dx = _cell_float(row[col[dx_key]])
        rcs = _cell_float(row[col[rcs_key]])
    except Exception:
        return None
    if math.isfinite(dx) and math.isfinite(rcs):
        return float(dx), float(rcs)
    return None


def build_filtered_rcs_rows_from_cluster_raw_path(
    raw_path: Path,
    calibration_db: float,
) -> Tuple[List[List[str]], List[Tuple[float, float]], int]:
    rows = _read_table_rows(raw_path)
    metadata_rows: List[List[str]] = []
    header_index = -1
    col: Dict[str, int] = {}
    for i, row in enumerate(rows):
        if row and _cell_text(row[0]) == "Time" and any(_cell_text(c) == "DX00" for c in row):
            header = [_cell_text(c) for c in row]
            col = {name: idx for idx, name in enumerate(header) if name}
            header_index = i
            metadata_rows = [list(r) for r in rows[:i]]
            break
    if header_index < 0:
        raise ValueError("未找到 Cluster Raw 表头行（需要包含 Time 和 DX00）")
    for required in ("R", "DX00", "RCS00"):
        if required not in col:
            raise ValueError(f"Cluster Raw 缺少列：{required}")

    max_ix = max(col.values())
    bins: Dict[int, List[float]] = defaultdict(list)
    selected_count = 0
    for raw_row in rows[header_index + 1 :]:
        if not raw_row or not any(_cell_text(c) for c in raw_row):
            continue
        row = list(raw_row)
        while len(row) <= max_ix:
            row.append("")
        try:
            range_val = _cell_float(row[col["R"]])
        except Exception:
            continue
        if not math.isfinite(range_val):
            continue

        candidates: List[Tuple[float, int, float, float]] = []
        for slot in range(MAX_CLUSTERS):
            parsed = _cluster_slot_dx_rcs(row, col, slot)
            if parsed is None:
                continue
            dx, rcs = parsed
            candidates.append((abs(float(dx) - float(range_val)), int(slot), float(dx), float(rcs)))
        if not candidates:
            continue
        candidates.sort(key=lambda item: (float(item[0]), int(item[1])))
        for _dist_err, _slot, dx, rcs in candidates[:2]:
            bin_id = int(round(float(dx) * 10.0))
            bins[bin_id].append(float(rcs) + float(calibration_db))
            selected_count += 1

    if not bins:
        raise ValueError("Cluster Raw 中没有可生成 Filtered 的有效帧")

    filtered_rows: List[Tuple[float, float]] = []
    for bin_id in range(min(bins.keys()), max(bins.keys()) + 1):
        values = bins.get(int(bin_id), [])
        x_val = float(bin_id) / 10.0
        if values:
            filtered_rows.append((x_val, float(sum(values)) / float(len(values))))
        else:
            filtered_rows.append((x_val, float("nan")))
    return metadata_rows, filtered_rows, selected_count


def filtered_rows_from_loaded_curve(curve: LoadedCurve, calibration_db: float) -> List[Tuple[float, float]]:
    try:
        _meta, rows, _count = build_filtered_rcs_rows_from_cluster_raw_path(curve.file_path, calibration_db)
        return rows
    except Exception:
        x, y = fit_segments(curve.segments)
        out: List[Tuple[float, float]] = []
        for x_val, y_val in zip(np.asarray(x, dtype=float).tolist(), np.asarray(y, dtype=float).tolist()):
            if math.isfinite(float(x_val)) and math.isfinite(float(y_val)):
                out.append((round(float(x_val), 1), float(y_val) + float(calibration_db)))
        return out


def write_rcs_two_column_csv(
    output_path: Path,
    columns: Tuple[str, str],
    rows: List[Tuple[float, float]],
    *,
    include_nan_rcs: bool,
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(list(columns))
        for x_val, rcs_val in rows:
            if not math.isfinite(float(x_val)):
                continue
            if not math.isfinite(float(rcs_val)):
                if not include_nan_rcs:
                    continue
                rcs_text = "NaN"
            else:
                rcs_text = f"{float(rcs_val):.4f}"
            writer.writerow([f"{float(x_val):.1f}", rcs_text])


def build_combined_rcs_rows_from_filtered_sets(
    filtered_sets: List[List[Tuple[float, float]]],
) -> List[Tuple[float, float]]:
    by_bin: Dict[int, List[float]] = defaultdict(list)
    for rows in filtered_sets:
        for x_val, rcs_val in rows:
            if math.isfinite(float(x_val)) and math.isfinite(float(rcs_val)):
                by_bin[int(round(float(x_val) * 10.0))].append(float(rcs_val))
    if not by_bin:
        return []
    xs: List[float] = []
    ys: List[float] = []
    for bin_id in sorted(by_bin.keys()):
        vals = by_bin[int(bin_id)]
        xs.append(float(bin_id) / 10.0)
        ys.append(float(sum(vals)) / float(len(vals)))
    return list(zip(xs, _smooth_rcs_series(np.asarray(ys, dtype=float)).tolist()))


def rows_to_bin_map(rows: List[Tuple[float, float]]) -> Dict[int, float]:
    out: Dict[int, float] = {}
    for x_val, y_val in rows:
        if math.isfinite(float(x_val)) and math.isfinite(float(y_val)):
            out[int(round(float(x_val) * 10.0))] = float(y_val)
    return out


def interpolate_curve_to_bins(xs: np.ndarray, ys: np.ndarray, bins: List[int]) -> Dict[int, float]:
    x = np.asarray(xs, dtype=float)
    y = np.asarray(ys, dtype=float)
    n = min(int(x.size), int(y.size))
    if n <= 0:
        return {}
    x = x[:n]
    y = y[:n]
    mask = np.isfinite(x) & np.isfinite(y)
    x = x[mask]
    y = y[mask]
    if x.size == 0:
        return {}
    order = np.argsort(x)
    x = x[order]
    y = y[order]
    out: Dict[int, float] = {}
    if x.size == 1:
        key = int(round(float(x[0]) * 10.0))
        if key in bins:
            out[key] = float(y[0])
        return out
    for bin_id in bins:
        d = float(bin_id) / 10.0
        if d < float(x[0]) or d > float(x[-1]):
            continue
        out[int(bin_id)] = float(np.interp(d, x, y))
    return out


def reference_limit_rows(limits: Optional[Dict[str, np.ndarray]]) -> List[Tuple[float, float, float]]:
    if not limits:
        return []
    xs = np.asarray(limits.get("x", []), dtype=float)
    lower = np.asarray(limits.get("lower", []), dtype=float)
    upper = np.asarray(limits.get("upper", []), dtype=float)
    n = min(int(xs.size), int(lower.size), int(upper.size))
    if n <= 0:
        return []
    xs = xs[:n]
    lower = lower[:n]
    upper = upper[:n]
    mask = np.isfinite(xs) & np.isfinite(lower) & np.isfinite(upper)
    xs = xs[mask]
    lower = lower[mask]
    upper = upper[mask]
    if xs.size == 0:
        return []
    order = np.argsort(xs)
    xs = xs[order]
    lower = lower[order]
    upper = upper[order]
    out: List[Tuple[float, float, float]] = []
    for d in np.arange(5.0, 50.0 + 1e-9, 5.0, dtype=float):
        if d < float(xs[0]) or d > float(xs[-1]):
            continue
        out.append((float(d), float(np.interp(d, xs, lower)), float(np.interp(d, xs, upper))))
    return out


class PlotCanvas(FigureCanvas):
    def __init__(self) -> None:
        self.figure = Figure(figsize=(8, 5), dpi=110)
        super().__init__(self.figure)
        self.setSizePolicy(QtWidgets.QSizePolicy.Expanding, QtWidgets.QSizePolicy.Expanding)


class MainWindow(QtWidgets.QMainWindow):
    def __init__(self) -> None:
        super().__init__()
        self.setWindowTitle(APP_TITLE)
        self.resize(1180, 760)
        self.curves: List[LoadedCurve] = []
        self.orbit_rows: Optional[List[Dict[str, float]]] = None
        self.point_cloud_points: List[PointCloudPoint] = []
        self.plot_mode = "distance"
        self.current_file_dir = Path.cwd()
        self._build_ui()
        self._populate_reference_options()
        self._redraw()

    def _build_ui(self) -> None:
        central = QtWidgets.QWidget()
        self.setCentralWidget(central)
        root = QtWidgets.QHBoxLayout(central)

        side = QtWidgets.QWidget()
        side.setMinimumWidth(320)
        side.setMaximumWidth(390)
        side_layout = QtWidgets.QVBoxLayout(side)

        form = QtWidgets.QFormLayout()
        self.product_combo = QtWidgets.QComboBox()
        self.angle_combo = QtWidgets.QComboBox()
        self.mode_combo = QtWidgets.QComboBox()
        self.mode_combo.addItem("直线距离-RCS", "distance")
        self.mode_combo.addItem("圆周 RCS", "orbit")
        self.mode_combo.addItem("空间点云图", "point_cloud")
        self.calibration_spin = QtWidgets.QDoubleSpinBox()
        self.calibration_spin.setRange(-80.0, 80.0)
        self.calibration_spin.setDecimals(2)
        self.calibration_spin.setSuffix(" dB")
        self.calibration_spin.setSingleStep(0.5)
        form.addRow("参考产品", self.product_combo)
        form.addRow("参考角度", self.angle_combo)
        form.addRow("绘图类型", self.mode_combo)
        form.addRow("RCS 标定值", self.calibration_spin)
        side_layout.addLayout(form)

        self.import_button = QtWidgets.QPushButton("导入文件")
        self.import_more_button = QtWidgets.QPushButton("导入对比文件")
        self.export_button = QtWidgets.QPushButton("保存图片")
        self.clear_button = QtWidgets.QPushButton("清空")
        grid = QtWidgets.QGridLayout()
        grid.addWidget(self.import_button, 0, 0)
        grid.addWidget(self.import_more_button, 0, 1)
        grid.addWidget(self.export_button, 1, 0)
        grid.addWidget(self.clear_button, 1, 1)
        side_layout.addLayout(grid)

        self.file_list = QtWidgets.QListWidget()
        self.file_list.setMinimumHeight(210)
        side_layout.addWidget(QtWidgets.QLabel("已导入数据"))
        side_layout.addWidget(self.file_list)

        self.status_text = QtWidgets.QTextEdit()
        self.status_text.setReadOnly(True)
        self.status_text.setMaximumHeight(180)
        side_layout.addWidget(QtWidgets.QLabel("状态"))
        side_layout.addWidget(self.status_text)
        side_layout.addStretch(1)

        plot_side = QtWidgets.QWidget()
        plot_layout = QtWidgets.QVBoxLayout(plot_side)
        self.canvas = PlotCanvas()
        self.toolbar = NavigationToolbar(self.canvas, self)
        plot_layout.addWidget(self.toolbar)
        plot_layout.addWidget(self.canvas, 1)

        root.addWidget(side)
        root.addWidget(plot_side, 1)

        self.import_button.clicked.connect(self.import_files)
        self.import_more_button.clicked.connect(self.import_compare_files)
        self.export_button.clicked.connect(self.export_image)
        self.clear_button.clicked.connect(self.clear_data)
        self.product_combo.currentIndexChanged.connect(self._redraw)
        self.angle_combo.currentIndexChanged.connect(self._redraw)
        self.mode_combo.currentIndexChanged.connect(self._mode_changed)
        self.calibration_spin.valueChanged.connect(self._redraw)

    def _populate_reference_options(self) -> None:
        classes, angles = rcs_reference_data.get_rcs_reference_options()
        labels = rcs_reference_data.get_rcs_reference_labels()
        self.product_combo.addItem("不叠加参考", "")
        for cls in classes:
            label = labels.get(cls, cls)
            self.product_combo.addItem(f"{label} ({cls})", cls)
        for angle in angles:
            self.angle_combo.addItem(angle, angle)

    def _mode_changed(self) -> None:
        self.plot_mode = str(self.mode_combo.currentData() or "distance")
        self._redraw()

    def _reference_limits(self) -> Optional[Dict[str, np.ndarray]]:
        cls = str(self.product_combo.currentData() or "")
        angle = str(self.angle_combo.currentData() or "")
        if not cls:
            return None
        return rcs_reference_data.get_rcs_reference_limits(cls, angle)

    def _append_status(self, text: str) -> None:
        self.status_text.append(text)

    def _choose_files(self, title: str) -> List[Path]:
        paths, _ = QtWidgets.QFileDialog.getOpenFileNames(
            self,
            title,
            str(self.current_file_dir),
            "RCS 表格 (*.csv *.xlsx *.xlsm);;CSV (*.csv);;Excel (*.xlsx *.xlsm);;所有文件 (*.*)",
        )
        if not paths:
            return []
        self.current_file_dir = Path(paths[0]).parent
        return [Path(p) for p in paths]

    def import_files(self) -> None:
        paths = self._choose_files("选择 RCS 数据文件")
        if not paths:
            return
        self.curves = []
        self.orbit_rows = None
        self.point_cloud_points = []
        self.file_list.clear()
        self._load_paths(paths, replace=True)

    def import_compare_files(self) -> None:
        paths = self._choose_files("选择要叠加对比的 RCS 数据文件")
        if not paths:
            return
        self._load_paths(paths, replace=False)

    def _load_paths(self, paths: List[Path], *, replace: bool) -> None:
        loaded = 0
        errors: List[str] = []
        for path in paths:
            try:
                if self.plot_mode == "point_cloud":
                    source_index = len({p.source_index for p in self.point_cloud_points}) + 1
                    start_idx = len(self.point_cloud_points) + 1
                    points, _ = parse_frame_aligned_point_cloud(
                        path,
                        source_index=source_index,
                        start_idx=start_idx,
                    )
                    self.point_cloud_points.extend(points)
                    self.curves = []
                    self.orbit_rows = None
                    self.file_list.addItem(f"{path.name} | 点云簇 {len(points)}")
                    loaded += 1
                    continue
                if self.plot_mode == "orbit":
                    if is_orbit_rcs_table(path):
                        rows = parse_orbit_rcs_table(path)
                    else:
                        segments, _ = parse_cluster_rcs_table(path)
                        rows = orbit_rows_from_cluster_segments(segments)
                    self.orbit_rows = rows
                    self.curves = []
                    self.file_list.clear()
                    self.file_list.addItem(f"{path.name} | 圆周点数 {len(rows)}")
                    loaded += 1
                    break
                curve = load_distance_curve(path)
                self.curves.append(curve)
                self.file_list.addItem(f"{path.name} | 点数 {curve.point_count}")
                loaded += 1
            except Exception as exc:
                errors.append(f"{path.name}: {exc}")

        if loaded:
            self._append_status(f"已导入 {loaded} 个文件")
        for err in errors:
            self._append_status(f"导入失败：{err}")
        self._redraw()

    def clear_data(self) -> None:
        self.curves = []
        self.orbit_rows = None
        self.point_cloud_points = []
        self.file_list.clear()
        self._append_status("已清空")
        self._redraw()

    def _export_distance_summary_files(self, image_path: Path) -> Optional[Path]:
        if self.plot_mode != "distance" or not self.curves:
            return None
        if openpyxl is None:
            raise RuntimeError("当前环境缺少 openpyxl，不能生成 Excel 汇总文件")

        calibration_db = float(self.calibration_spin.value())
        if len(self.curves) > 1:
            ava_x, ava_y = fit_loaded_curves_combined(self.curves)
        else:
            ava_x, ava_y = self.curves[0].fitted
        ava_y = np.asarray(ava_y, dtype=float) + calibration_db

        curve_xy: List[Tuple[np.ndarray, np.ndarray]] = []
        for curve in self.curves:
            x_vals = np.asarray(curve.fitted[0], dtype=float)
            y_vals = np.asarray(curve.fitted[1], dtype=float) + calibration_db
            curve_xy.append((x_vals, y_vals))

        bin_ids: set[int] = set()
        for x_vals, y_vals in curve_xy:
            x = np.asarray(x_vals, dtype=float)
            y = np.asarray(y_vals, dtype=float)
            n = min(int(x.size), int(y.size))
            if n <= 0:
                continue
            x = x[:n]
            y = y[:n]
            mask = np.isfinite(x) & np.isfinite(y) & (x <= 50.0)
            if not np.any(mask):
                continue
            lo = int(math.ceil(float(np.min(x[mask])) * 10.0 - 1e-9))
            hi = int(math.floor(float(np.max(x[mask])) * 10.0 + 1e-9))
            for bin_id in range(lo, min(hi, 500) + 1):
                bin_ids.add(int(bin_id))

        all_bins = sorted(bin_ids)
        rcs_maps = [
            interpolate_curve_to_bins(x_vals, y_vals, all_bins)
            for x_vals, y_vals in curve_xy
        ]
        ava_map = interpolate_curve_to_bins(np.asarray(ava_x, dtype=float), ava_y, all_bins)
        all_bins = sorted(bin_id for bin_id in (set(all_bins) | set(ava_map.keys())) if int(bin_id) <= 500)

        report_path = image_path.with_name(f"{image_path.stem}_summary.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Summary"

        headers = ["距离"] + [f"RCS{i + 1}" for i in range(len(rcs_maps))] + ["RCSAVA"]
        ws.append(headers)
        for bin_id in all_bins:
            row = [float(bin_id) / 10.0]
            for m in rcs_maps:
                row.append(m.get(int(bin_id)))
            row.append(ava_map.get(int(bin_id)))
            ws.append(row)

        header_fill = PatternFill("solid", fgColor="F2F2F2")
        thin = Side(style="thin", color="BFBFBF")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = border
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
            for idx, cell in enumerate(row, start=1):
                cell.border = border
                cell.number_format = "0.0" if idx == 1 else "0.0000"
        ws.freeze_panes = "A2"
        ws.column_dimensions["A"].width = 10
        for col_idx in range(2, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = 13

        image_for_excel = image_path
        temp_image_path: Optional[Path] = None
        if image_path.suffix.lower() not in {".png", ".jpg", ".jpeg"}:
            tmp = tempfile.NamedTemporaryFile(prefix="rcs_summary_chart_", suffix=".png", delete=False)
            tmp.close()
            temp_image_path = Path(tmp.name)
            image_for_excel = temp_image_path
            self.canvas.figure.savefig(image_for_excel, dpi=220, bbox_inches="tight")
        if image_for_excel.exists():
            img = OpenpyxlImage(str(image_for_excel))
            img.width = 500
            img.height = 300
            ws.add_image(img, "G2")

        ref_rows = reference_limit_rows(self._reference_limits())
        if ref_rows:
            start_row = 23
            start_col = 8
            ref_headers = ["距离(m)", "下限(dBsm)", "上限(dBsm)"]
            for offset, header in enumerate(ref_headers):
                cell = ws.cell(row=start_row, column=start_col + offset, value=header)
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center")
                cell.border = border
            for r_offset, (dist, lower, upper) in enumerate(ref_rows, start=1):
                values = [dist, lower, upper]
                for c_offset, value in enumerate(values):
                    cell = ws.cell(row=start_row + r_offset, column=start_col + c_offset, value=value)
                    cell.border = border
                    cell.number_format = "0.0"
            for c_offset in range(3):
                ws.column_dimensions[get_column_letter(start_col + c_offset)].width = 13

        wb.save(report_path)
        if temp_image_path is not None:
            try:
                temp_image_path.unlink(missing_ok=True)
            except Exception:
                pass
        self._append_status(f"汇总文件已生成：{report_path.name}")
        return report_path

    def export_image(self) -> None:
        if self.plot_mode == "orbit":
            default_name = "rcs_orbit.png"
        elif self.plot_mode == "point_cloud":
            default_name = "rcs_point_cloud.png"
        else:
            default_name = "rcs_distance.png"
        path, _ = QtWidgets.QFileDialog.getSaveFileName(
            self,
            "保存图片",
            str(self.current_file_dir / default_name),
            "PNG 图片 (*.png);;JPEG 图片 (*.jpg);;PDF 文件 (*.pdf);;所有文件 (*.*)",
        )
        if not path:
            return
        try:
            image_path = Path(path)
            self.canvas.figure.savefig(image_path, dpi=320, bbox_inches="tight")
            self._append_status(f"图片已保存：{image_path}")
            if self.plot_mode == "distance" and self.curves:
                report_path = self._export_distance_summary_files(image_path)
                if report_path is not None:
                    self._append_status(f"Excel汇总已保存：{report_path}")
        except Exception as exc:
            QtWidgets.QMessageBox.warning(self, "保存失败", str(exc))

    def _redraw(self) -> None:
        if not hasattr(self, "canvas"):
            return
        self.canvas.figure.clear()
        if self.plot_mode == "orbit":
            ax = self.canvas.figure.add_subplot(111, projection="polar")
            self._draw_orbit(ax)
        elif self.plot_mode == "point_cloud":
            ax = self.canvas.figure.add_subplot(111)
            self._draw_point_cloud(ax)
        else:
            ax = self.canvas.figure.add_subplot(111)
            self._draw_distance(ax)
        self.canvas.figure.tight_layout()
        self.canvas.draw_idle()

    def _draw_distance(self, ax) -> None:
        ax.clear()
        ax.set_title("RCS 距离曲线")
        ax.set_xlabel("Front Distance (m)")
        ax.set_ylabel("RCS (dBsm)")
        ax.set_xlim(0.0, RCS_MAX_DISTANCE_M)
        ax.grid(True, axis="y", color="#DCE3E8", linewidth=0.8)
        ax.grid(True, axis="x", color="#EEF2F5", linestyle="--", linewidth=0.65)
        ax.spines["top"].set_visible(False)
        ax.spines["right"].set_visible(False)

        colors = ["#1565C0", "#FFB74D", "#AD1457", "#6A1B9A", "#EF6C00", "#00838F", "#5D4037", "#283593"]
        cal = float(self.calibration_spin.value())
        for i, curve in enumerate(self.curves):
            color = colors[i % len(colors)]
            if len(self.curves) == 1:
                for j, seg in enumerate(curve.segments[:30]):
                    xb, yb = single_measurement_curve(seg)
                    if xb.size:
                        ax.plot(xb, yb + cal, color=colors[j % len(colors)], linewidth=1.15, alpha=0.82, label=f"第{j + 1}次测量" if len(curve.segments) > 1 else curve.display_name)
            xg, yg = curve.fitted
            if xg.size and yg.size:
                label = curve.display_name if len(self.curves) > 1 else ("融合曲线" if len(curve.segments) > 1 else None)
                ax.plot(xg, yg + cal, color=color if len(self.curves) > 1 else "#2E7D32", linewidth=2.1, label=label)

        if len(self.curves) > 1:
            xg, yg = fit_loaded_curves_combined(self.curves)
            if xg.size and yg.size:
                ax.plot(
                    xg,
                    yg + cal,
                    color="#2E7D32",
                    linewidth=2.7,
                    alpha=0.98,
                    label="combined",
                    zorder=5,
                )

        limits = self._reference_limits()
        if limits is not None:
            xs = np.asarray(limits.get("x"), dtype=float)
            mask = np.isfinite(xs) & (xs >= 0.0) & (xs <= RCS_MAX_DISTANCE_M)
            lower = limits.get("lower")
            upper = limits.get("upper")
            if lower is not None:
                lower = np.asarray(lower, dtype=float)
                if lower.shape == xs.shape:
                    ax.plot(xs[mask], lower[mask], color="black", linewidth=1.7, label="参考下限")
            if upper is not None:
                upper = np.asarray(upper, dtype=float)
                if upper.shape == xs.shape:
                    ax.plot(xs[mask], upper[mask], color="black", linewidth=1.7, label="参考上限")

        handles, labels = ax.get_legend_handles_labels()
        if labels:
            by_label = dict(zip(labels, handles))
            ax.legend(by_label.values(), by_label.keys(), loc="best", frameon=False)
        if not self.curves:
            ax.text(0.5, 0.5, "请导入 Cluster RCS CSV/Excel 表格", transform=ax.transAxes, ha="center", va="center", color="#607D8B")

    def _draw_point_cloud(self, ax) -> None:
        ax.clear()
        ax.set_title("空间点云图")
        ax.set_xlabel("Y, m")
        ax.set_ylabel("X, m")
        ax.set_xlim(-1.7, 1.7)
        ax.set_ylim(-2.2, 2.2)
        ax.grid(True, color="#DCE3E8", linewidth=0.8, alpha=0.75)
        ax.set_aspect("equal", adjustable="box")

        visible = [
            point
            for point in self.point_cloud_points
            if math.isfinite(point.x)
            and math.isfinite(point.y)
            and math.isfinite(point.rcs_raw)
            and abs(point.x) <= 2.2
            and abs(point.y) <= 1.7
        ]
        if not visible:
            ax.text(
                0.5,
                0.5,
                "请选择“空间点云图”后导入 Cluster Raw 文件",
                transform=ax.transAxes,
                ha="center",
                va="center",
                color="#607D8B",
            )
            return

        calibration_db = float(self.calibration_spin.value())
        scatter = ax.scatter(
            [point.y for point in visible],
            [point.x for point in visible],
            c=[point.rcs_raw + calibration_db for point in visible],
            s=7,
            marker="s",
            cmap="jet",
            linewidths=0,
            alpha=0.95,
        )
        colorbar = self.canvas.figure.colorbar(scatter, ax=ax, pad=0.02, fraction=0.045)
        colorbar.set_label("RCS, dBsm")
        ax.text(
            0.01,
            0.99,
            f"文件 {len({point.source_index for point in self.point_cloud_points})} | "
            f"有效簇 {len(self.point_cloud_points)} | 窗口内 {len(visible)}",
            transform=ax.transAxes,
            ha="left",
            va="top",
            fontsize=9,
            color="#455A64",
        )

    def _draw_orbit(self, ax) -> None:
        ax.clear()
        ax.set_title("圆周 RCS")
        ax.set_theta_zero_location("N")
        ax.set_theta_direction(-1)
        theta_ticks = np.deg2rad(np.arange(0, 360, 30))
        ax.set_xticks(theta_ticks)
        ax.set_xticklabels([f"{deg}°" for deg in np.arange(0, 360, 30)], fontsize=11, fontweight="bold")
        ax.grid(True, linestyle="--", alpha=0.5, linewidth=0.8)
        if not self.orbit_rows:
            ax.text(0.5, 0.5, "请切换到圆周 RCS 后导入圆周表格或 Cluster Raw", transform=ax.transAxes, ha="center", va="center", color="#607D8B")
            return

        series = build_orbit_series(self.orbit_rows)
        angles = np.asarray(series["angles"], dtype=float)
        values = np.asarray(series["values"], dtype=float) + float(self.calibration_spin.value())
        angles, values = smooth_orbit_line(angles, values)
        if angles.size == 0 or values.size == 0:
            ax.text(0.5, 0.5, "圆周 RCS 有效点不足", transform=ax.transAxes, ha="center", va="center", color="#607D8B")
            return

        angles_closed = np.append(angles, angles[0])
        values_closed = np.append(values, values[0])
        ax.plot(angles_closed, values_closed, color="#4472C4", linewidth=1.2, alpha=0.92)

        finite_values = values[np.isfinite(values)]
        if finite_values.size:
            data_min = float(np.min(finite_values))
            data_max = float(np.max(finite_values))
        else:
            data_min, data_max = -15.0, 5.0
        r_min = min(-16.0, math.floor((data_min - 1.0) / 5.0) * 5.0)
        r_max = max(6.0, math.ceil((data_max + 1.0) / 5.0) * 5.0)
        if r_max <= r_min:
            r_max = r_min + 5.0
        tick_start = math.ceil(r_min / 5.0) * 5.0
        tick_end = math.floor(r_max / 5.0) * 5.0
        r_ticks = np.arange(tick_start, tick_end + 0.1, 5.0)
        ax.set_ylim(r_min, r_max)
        ax.set_yticks(r_ticks)
        ax.set_yticklabels([f"{int(tick)}" for tick in r_ticks], fontsize=10, fontweight="bold")
        ax.text(0.5, 0.95, "dBsm", transform=ax.transAxes, fontsize=11, fontweight="bold", ha="center", va="center")


def main() -> int:
    _apply_matplotlib_font()
    app = QtWidgets.QApplication(sys.argv)
    app.setApplicationName(APP_TITLE)
    window = MainWindow()
    window.show()
    return int(app.exec_())


if __name__ == "__main__":
    raise SystemExit(main())
